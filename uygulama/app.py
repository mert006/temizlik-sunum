# -*- coding: utf-8 -*-
"""Vega Taurus AVM — Temizlik Birimi Aylık Sunum Sihirbazı.

Bu dosya başlatıcı (Temizlik Sunum.exe) tarafından her açılışta yüklenir.
İnternetten güncelleme gelirse bu dosyanın yerine yenisi yazılır.
"""
import copy
import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import traceback
import uuid

import tkinter as tk
from tkinter import ttk, messagebox, filedialog

from PIL import Image, ImageOps, ImageTk
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

AYLAR = ["OCAK", "ŞUBAT", "MART", "NİSAN", "MAYIS", "HAZİRAN", "TEMMUZ",
         "AĞUSTOS", "EYLÜL", "EKİM", "KASIM", "ARALIK"]
AYLAR_KUCUK = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz",
               "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]

MAVI = RGBColor(0x44, 0x72, 0xC4)
KOYU = RGBColor(0x26, 0x26, 0x26)
KIRMIZI = RGBColor(0xC0, 0x00, 0x00)

UYGULAMA_DIR = os.path.dirname(os.path.abspath(__file__))


# ---------------------------------------------------------------------------
# Türkçe metin yardımcıları
# ---------------------------------------------------------------------------

def tr_lower(s):
    return s.replace("I", "ı").replace("İ", "i").lower()


def tr_upper(s):
    return s.replace("i", "İ").replace("ı", "I").upper()


def ilk_harf_buyuk(s):
    for i, ch in enumerate(s):
        if ch.isalpha():
            return s[:i] + tr_upper(ch) + s[i + 1:]
    return s


def ayarlari_yukle():
    with open(os.path.join(UYGULAMA_DIR, "ayarlar.json"), encoding="utf-8") as f:
        return json.load(f)


def metni_duzelt(metin, ayar):
    """Yazım biçimini standarda çeker. (düzeltilmiş_metin, yapılan_değişiklikler) döner."""
    degisiklik = []
    t = " ".join(metin.replace("\n", " ").split())
    if not t:
        return t, degisiklik

    harfler = [c for c in t if c.isalpha()]
    if len(harfler) > 8 and sum(1 for c in harfler if c.isupper()) / len(harfler) > 0.6:
        t = tr_lower(t)
        degisiklik.append("TAMAMI BÜYÜK HARF yazılmıştı, normal yazıma çevrildi.")

    eski = t
    t = re.sub(r"\s+([,.;:!?])", r"\1", t)                       # noktalama öncesi boşluk
    t = re.sub(r"(?<!\d),(?!\d)\s*", ", ", t)                     # virgülden sonra boşluk
    t = re.sub(r";\s*", "; ", t)
    t = re.sub(r"\.(?=[A-Za-zÇĞİÖŞÜçğıöşü])", ". ", t)             # noktadan sonra boşluk
    t = re.sub(r"\s{2,}", " ", t).strip()
    t = re.sub(r"[.]{2,}$", ".", t)
    if t and t[-1] not in ".!?":
        t += "."
    if t != eski:
        degisiklik.append("Boşluk/noktalama düzeltildi.")

    for yanlis, dogru in ayar.get("yazim_duzeltmeleri", {}).items():
        desen = r"(?<!\w)" + re.escape(yanlis) + r"(?!\w)"
        if re.search(desen, t, flags=re.IGNORECASE):
            t = re.sub(desen, dogru, t, flags=re.IGNORECASE)
            degisiklik.append(f"Yazım: '{yanlis}' → '{dogru}'")

    for isim in ayar.get("ozel_isimler", []):
        desen = r"(?<!\w)" + re.escape(tr_lower(isim)) + r"(?!\w)"
        dusuk = tr_lower(t)
        for m in reversed(list(re.finditer(desen, dusuk))):
            t = t[:m.start()] + isim + t[m.end():]

    # cümle başları büyük harf
    parcalar = re.split(r"((?<!\d)[.!?]\s+)", t)  # "2. kat" gibi sıra sayılarında bölme
    t = "".join(ilk_harf_buyuk(p) if i % 2 == 0 else p for i, p in enumerate(parcalar))
    return t, degisiklik


def metin_uyarilari(metin, ayar):
    """Engellemeyen (yumuşak) uyarılar."""
    uyari = []
    kelime = len(metin.split())
    if kelime < ayar.get("min_kelime", 8):
        uyari.append(f"Cümle kısa ({kelime} kelime). Nerede (kat/blok/alan), ne yapıldığını ve kapsamını yazın.")
    dusuk = tr_lower(metin)
    if not any(k in dusuk for k in ayar.get("konum_kelimeleri", [])):
        uyari.append("Konum yazılmamış görünüyor (ör. '2. kat', 'P2 otopark', 'Konya Yolu girişi').")
    for d in ayar.get("dolgu_ifadeler", []):
        if tr_lower(d) in dusuk:
            uyari.append(f"'{d}' kalıp bir ifade. Bunun yerine somut bilgi yazın (kaç adet, hangi alan).")
    if not re.search(r"\d", metin):
        uyari.append("İpucu: Mümkünse sayı ekleyin (ör. '12 adet cam', '3 kat', '2 gün').")
    return uyari


# ---------------------------------------------------------------------------
# Fotoğraf
# ---------------------------------------------------------------------------

def fotografi_hazirla(kaynak, hedef_klasor):
    """Fotoğrafı döndürür, küçültür, kayıt klasörüne kopyalar. {'dosya','w','h'} döner."""
    im = Image.open(kaynak)
    im = ImageOps.exif_transpose(im)
    im = im.convert("RGB")
    im.thumbnail((1920, 1920))
    os.makedirs(hedef_klasor, exist_ok=True)
    ad = uuid.uuid4().hex[:10] + ".jpg"
    im.save(os.path.join(hedef_klasor, ad), "JPEG", quality=85)
    return {"dosya": ad, "w": im.width, "h": im.height}


# ---------------------------------------------------------------------------
# Sunum üretimi
# ---------------------------------------------------------------------------

def _layout(prs, ad):
    for l in prs.slide_layouts:
        if l.name == ad:
            return l
    raise RuntimeError(f"Şablonda '{ad}' düzeni yok")


def _yer_tutuculari_sil(slide, haric=()):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx not in haric:
            ph._element.getparent().remove(ph._element)


def _altyazi_yaz(ph, metin):
    tf = ph.text_frame
    tf.text = metin
    tf.word_wrap = True
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
            for e in pPr.findall(qn(tag)):
                pPr.remove(e)
        pPr.insert(0, pPr.makeelement(qn("a:buNone"), {}))
        for r in p.runs:
            r.font.size = Pt(24)
            r.font.bold = False


def _hucre(cell, metin, renk=None, kalin=None, boyut=None):
    tf = cell.text_frame
    paras = tf.paragraphs
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    p0 = paras[0]
    runs = p0.runs
    if runs:
        runs[0].text = str(metin)
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        run = runs[0]
    else:
        run = p0.add_run()
        run.text = str(metin)
    if renk is not None:
        run.font.color.rgb = renk
    if kalin is not None:
        run.font.bold = kalin
    if boyut is not None:
        run.font.size = Pt(boyut)


def _metin_kutusu(slide, x, y, w, h, metin, boyut, kalin=False, renk=KOYU, hiza=PP_ALIGN.LEFT, dikey=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = dikey
    satirlar = metin.split("\n")
    for i, s in enumerate(satirlar):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = hiza
        r = p.add_run()
        r.text = s
        r.font.size = Pt(boyut)
        r.font.bold = kalin
        r.font.color.rgb = renk
    return tb


def _foto_yolu(kayit_dir, f):
    return os.path.join(kayit_dir, "foto", f["dosya"])


def _is_slayti(prs, metin, fotolar, kayit_dir):
    n = len(fotolar)
    dikey = [f["h"] > f["w"] for f in fotolar]
    if n == 1:
        s = prs.slides.add_slide(_layout(prs, "Başlık ve İçerik"))
        _yer_tutuculari_sil(s, haric=(11,))
        f = fotolar[0]
        kx, ky, kw, kh = 1.58, 0.40, 10.17, 5.07
        oran = min(kw / f["w"], kh / f["h"])
        w, h = f["w"] * oran, f["h"] * oran
        s.shapes.add_picture(_foto_yolu(kayit_dir, f), Inches(kx + (kw - w) / 2), Inches(ky + (kh - h) / 2),
                             Inches(w), Inches(h))
        _altyazi_yaz(s.placeholders[11], metin)
        return s
    if n == 2:
        s = prs.slides.add_slide(_layout(prs, "2_Boş"))
        idx = [13, 17]
    elif all(dikey):
        s = prs.slides.add_slide(_layout(prs, "1_Boş"))
        idx = [13, 14, 15]
    else:
        s = prs.slides.add_slide(_layout(prs, "3_Boş"))
        idx = [13, 17, 18]
    for i, f in zip(idx, fotolar):
        s.placeholders[i].insert_picture(_foto_yolu(kayit_dir, f))
    _altyazi_yaz(s.placeholders[16], metin)
    return s


def genel_durum_cumlesi(v):
    ay = AYLAR_KUCUK[v["ay"] - 1]
    fiili = sum(v["kadro"]["fiili"])
    gerc = sum(v["kadro"]["gerceklesen"])
    eksik = fiili - gerc
    p = v["performans"]
    c = [f"{ay} ayında departmanımız {fiili} kişilik tanımlı kadroya karşılık {gerc} personelle hizmet vermiştir"]
    c[0] += f" ({eksik} eksik personel)." if eksik > 0 else "."
    s, o = p["sikayet"], p.get("onceki_sikayet")
    if o is None:
        c.append(f"Ay içinde {s} şikayet alınmıştır.")
    elif o == 0:
        c.append(f"Ay içinde {s} şikayet alınmıştır (geçen ay şikayet alınmamıştı).")
    else:
        fark = (s - o) / o * 100
        if abs(fark) < 0.5:
            c.append(f"Ay içinde {s} şikayet alınmış olup geçen ayla aynı seviyededir.")
        else:
            yon = "azalmıştır" if fark < 0 else "artmıştır"
            c.append(f"Ay içinde {s} şikayet alınmış olup geçen aya göre %{abs(fark):.0f} {yon}.")
    d, u = p["denetim"], p["uygunsuzluk"]
    if d:
        if u:
            c.append(f"{d} hijyen denetimi yapılmış, tespit edilen {u} uygunsuzluk giderilmiştir.")
        else:
            c.append(f"{d} hijyen denetimi yapılmış, uygunsuzluk tespit edilmemiştir.")
    if p.get("egitim"):
        c.append(f"{p['egitim']} personele eğitim verilmiştir.")
    ek = (p.get("ek_not") or "").strip()
    if ek:
        c.append(ek)
    return " ".join(c)


def _genel_durum_slayti(prs, v):
    s = prs.slides.add_slide(_layout(prs, "Boş"))
    _yer_tutuculari_sil(s)
    ay = AYLAR_KUCUK[v["ay"] - 1]
    _metin_kutusu(s, 0.5, 0.15, 12.33, 0.8, f"Temizlik — {ay} Ayı Genel Durum Değerlendirmesi", 28, kalin=True,
                  hiza=PP_ALIGN.CENTER)
    fiili = sum(v["kadro"]["fiili"])
    gerc = sum(v["kadro"]["gerceklesen"])
    p = v["performans"]
    s_alt = "Bu ay alınan şikayet"
    o = p.get("onceki_sikayet")
    if o:
        fark = (p["sikayet"] - o) / o * 100
        s_alt = f"Şikayet (geçen ay {o}, %{abs(fark):.0f} {'azalış' if fark < 0 else 'artış' if fark > 0 else 'değişim'})"
    elif o == 0:
        s_alt = "Şikayet (geçen ay 0)"
    kartlar = [
        (str(fiili), "Tanımlı kadro"),
        (str(gerc), f"Fiilen çalışan personel\n({fiili - gerc} eksik)" if fiili > gerc else "Fiilen çalışan personel"),
        (str(p["sikayet"]), s_alt),
        (str(p["denetim"]), f"Hijyen denetimi\n({p['uygunsuzluk']} uygunsuzluk)"),
    ]
    kw, gap, top, kh = 2.75, 0.3, 1.25, 2.1
    x0 = (13.333 - (4 * kw + 3 * gap)) / 2
    for i, (sayi, etiket) in enumerate(kartlar):
        x = x0 + i * (kw + gap)
        kart = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(top), Inches(kw), Inches(kh))
        kart.adjustments[0] = 0.08
        kart.fill.solid()
        kart.fill.fore_color.rgb = MAVI
        kart.line.fill.background()
        kart.shadow.inherit = False
        _metin_kutusu(s, x, top + 0.15, kw, 1.0, sayi, 48, kalin=True, renk=RGBColor(0xFF, 0xFF, 0xFF),
                      hiza=PP_ALIGN.CENTER)
        _metin_kutusu(s, x + 0.1, top + 1.15, kw - 0.2, 0.9, etiket, 14, renk=RGBColor(0xFF, 0xFF, 0xFF),
                      hiza=PP_ALIGN.CENTER)
    _metin_kutusu(s, 0.9, 3.75, 11.55, 2.4, genel_durum_cumlesi(v), 20, hiza=PP_ALIGN.LEFT)
    return s


def _kadro_slaytini_doldur(slide, v):
    tablolar = [sh for sh in slide.shapes if sh.has_table]
    tablolar.sort(key=lambda sh: sh.top)
    t1, t2 = tablolar[0].table, tablolar[1].table
    for r, anahtar in ((1, "fiili"), (2, "gerceklesen")):
        degerler = v["kadro"][anahtar]
        for c, d in enumerate(degerler, start=1):
            _hucre(t1.cell(r, c), d)
        _hucre(t1.cell(r, len(degerler) + 1), sum(degerler))
    satir = 1
    for i, vd in enumerate(v["vardiya"], start=1):
        for anahtar, etiket in (("plan", f"{i}.Vardiya"), ("gercek", f"{i}.Vardiya Gerçekleşen")):
            _hucre(t2.cell(satir, 0), etiket)
            for c, d in enumerate(vd[anahtar], start=1):
                _hucre(t2.cell(satir, c), d)
            _hucre(t2.cell(satir, len(vd[anahtar]) + 1), sum(vd[anahtar]))
            satir += 1


def _makine_slaytini_doldur(slide, makineler):
    sh = [x for x in slide.shapes if x.has_table][0]
    tbl = sh.table._tbl
    trs = tbl.tr_lst
    ornek = copy.deepcopy(trs[1])
    for tr in trs[1:]:
        tbl.remove(tr)
    toplam_h = sh.height
    n = len(makineler)
    satir_h = int(toplam_h / (n + 1)) if n > 13 else int(trs[0].h)
    trs[0].h = satir_h
    for _ in makineler:
        yeni = copy.deepcopy(ornek)
        yeni.h = satir_h
        tbl.append(yeni)
    boyut = 12 if n > 13 else None
    for i, (marka, model, ozellik, durum) in enumerate(makineler, start=1):
        cells = sh.table.rows[i].cells
        _hucre(cells[0], i, boyut=boyut)
        _hucre(cells[1], marka, boyut=boyut)
        _hucre(cells[2], model or "-", boyut=boyut)
        _hucre(cells[3], ozellik, boyut=boyut)
        ariza = durum != "AKTİF"
        _hucre(cells[4], durum, renk=KIRMIZI if ariza else KOYU, kalin=ariza, boyut=boyut)
    sh.height = Emu(satir_h * (n + 1))


def _makine_dairesi_slayti(prs, v, kayit_dir):
    fotolar = v["makine_dairesi"]["fotolar"]
    s = prs.slides.add_slide(_layout(prs, "2_Boş"))
    _metin_kutusu(s, 0, 0.2, 13.333, 0.8, "Makine dairesi", 32, hiza=PP_ALIGN.CENTER)
    phs = {ph.placeholder_format.idx: ph for ph in s.placeholders}
    if len(fotolar) == 1:
        phs[17]._element.getparent().remove(phs[17]._element)
        solar = [(13.333 - 4.43) / 2]
    else:
        solar = [1.54, 7.03]
    for idx, f, sol in zip((13, 17), fotolar, solar):
        pic = phs[idx].insert_picture(_foto_yolu(kayit_dir, f))
        # konum/boyut dördü birden verilmeli, yoksa boyut 0 olur
        pic.left, pic.top, pic.width, pic.height = Inches(sol), Inches(1.23), Inches(4.43), Inches(4.62)
    not_ = (v["makine_dairesi"].get("not") or "").strip()
    if not_:
        _altyazi_yaz(phs[16], not_)
        ph = phs[16]
        ph.left, ph.top, ph.width, ph.height = Inches(1.17), Inches(6.0), Inches(10.99), Inches(0.6)
    else:
        phs[16]._element.getparent().remove(phs[16]._element)
    return s


def sunumu_olustur(v, kayit_dir, cikti_yolu):
    prs = Presentation(os.path.join(UYGULAMA_DIR, "sablon.pptx"))
    lst = prs.slides._sldIdLst
    ayrac, kadro, parkur = list(lst)
    sablon_slaytlar = list(prs.slides)
    _kadro_slaytini_doldur(sablon_slaytlar[1], v)
    _makine_slaytini_doldur(sablon_slaytlar[2], v["makineler"])

    def ekle(fonk, *a):
        fonk(prs, *a)
        return list(lst)[-1]

    sira = [ayrac, ekle(_genel_durum_slayti, v), kadro]
    if v["makine_dairesi"]["fotolar"]:
        sira.append(ekle(_makine_dairesi_slayti, v, kayit_dir))
    sira.append(parkur)
    for is_ in v["isler"] + v["etkinlikler"]:
        sira.append(ekle(_is_slayti, is_["metin"], is_["fotolar"], kayit_dir))
    for e in sira:
        lst.remove(e)
        lst.append(e)
    prs.save(cikti_yolu)
    return len(sira)


# ---------------------------------------------------------------------------
# Kayıt (taslak) yönetimi
# ---------------------------------------------------------------------------

def bos_veri(ayar, yil, ay):
    kr = len(ayar["kadro_rolleri"])
    vr = len(ayar["vardiya_rolleri"])
    return {
        "yil": yil, "ay": ay, "hazirlayan": "",
        "kadro": {"fiili": [0] * kr, "gerceklesen": [0] * kr},
        "vardiya": [{"plan": [0] * vr, "gercek": [0] * vr} for _ in range(ayar["vardiya_sayisi"])],
        "performans": {"sikayet": 0, "onceki_sikayet": None, "denetim": 0, "uygunsuzluk": 0, "egitim": 0,
                       "ek_not": ""},
        "makine_dairesi": {"fotolar": [], "not": ""},
        "makineler": [list(m) for m in ayar["varsayilan_makineler"]],
        "isler": [], "etkinlikler": [],
        "tamamlandi": False,
    }


class Kayitlar:
    def __init__(self, kok):
        self.kok = os.path.join(kok, "kayitlar")
        os.makedirs(self.kok, exist_ok=True)

    def klasor(self, yil, ay):
        return os.path.join(self.kok, f"{yil}-{ay:02d}")

    def oku(self, yil, ay):
        yol = os.path.join(self.klasor(yil, ay), "veri.json")
        if os.path.exists(yol):
            with open(yol, encoding="utf-8") as f:
                return json.load(f)
        return None

    def yaz(self, v):
        d = self.klasor(v["yil"], v["ay"])
        os.makedirs(d, exist_ok=True)
        gecici = os.path.join(d, "veri.json.tmp")
        with open(gecici, "w", encoding="utf-8") as f:
            json.dump(v, f, ensure_ascii=False, indent=1)
        os.replace(gecici, os.path.join(d, "veri.json"))

    def onceki_ay(self, yil, ay):
        hedef = f"{yil}-{ay:02d}"
        adaylar = sorted(x for x in os.listdir(self.kok) if re.fullmatch(r"\d{4}-\d{2}", x) and x < hedef)
        for a in reversed(adaylar):
            y, m = map(int, a.split("-"))
            v = self.oku(y, m)
            if v:
                return v
        return None


def masaustu():
    try:
        import winreg
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                            r"Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders") as k:
            yol = os.path.expandvars(winreg.QueryValueEx(k, "Desktop")[0])
            if os.path.isdir(yol):
                return yol
    except Exception:
        pass
    return os.path.join(os.path.expanduser("~"), "Desktop")


# ---------------------------------------------------------------------------
# Arayüz
# ---------------------------------------------------------------------------

YAZI = ("Segoe UI", 11)
YAZI_K = ("Segoe UI", 11, "bold")
BASLIK = ("Segoe UI", 18, "bold")
ACIKLAMA = ("Segoe UI", 10)
TURUNCU = "#b85c00"
YESIL = "#1e7b34"
GRI = "#666666"


class Kaydirmali(ttk.Frame):
    def __init__(self, parent, **kw):
        super().__init__(parent, **kw)
        self.canvas = tk.Canvas(self, highlightthickness=0, bg="#ffffff")
        sb = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.ic = ttk.Frame(self.canvas, style="Beyaz.TFrame")
        self.ic.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self._pencere = self.canvas.create_window((0, 0), window=self.ic, anchor="nw")
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfigure(self._pencere, width=e.width))
        self.canvas.configure(yscrollcommand=sb.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        self.canvas.bind("<Enter>", lambda e: self.canvas.bind_all("<MouseWheel>", self._tekerlek))
        self.canvas.bind("<Leave>", lambda e: self.canvas.unbind_all("<MouseWheel>"))

    def _tekerlek(self, e):
        self.canvas.yview_scroll(int(-e.delta / 120), "units")


def sayi_kutusu(parent, deger, genislik=6, degisince=None):
    var = tk.StringVar(value="" if deger is None else str(deger))
    e = ttk.Entry(parent, textvariable=var, width=genislik, justify="center", font=YAZI)
    if degisince:
        var.trace_add("write", lambda *a: degisince())
    return e, var


def sayi_oku(var, ad):
    t = var.get().strip()
    if t == "":
        raise ValueError(f"'{ad}' boş bırakılmış. Yoksa 0 yazın.")
    if not t.isdigit():
        raise ValueError(f"'{ad}' alanına sadece sayı yazın (yazılan: '{t}').")
    return int(t)


class Uygulama:
    def __init__(self, kok, surum, guncelleme_notu=None):
        self.kok = kok
        self.surum = surum
        self.ayar = ayarlari_yukle()
        self.kayit = Kayitlar(kok)
        self.v = None
        self.adim = 0
        self.thumbs = []
        self.root = tk.Tk()
        self.root.title(f"Vega Taurus AVM — Temizlik Sunum Sihirbazı (sürüm {surum})")
        self.root.geometry("1180x780")
        self.root.minsize(1000, 640)
        try:
            self.root.state("zoomed")
        except tk.TclError:
            pass
        self.root.configure(bg="#ffffff")
        self.root.option_add("*Font", YAZI)
        st = ttk.Style()
        try:
            st.theme_use("vista")
        except tk.TclError:
            pass
        st.configure("Beyaz.TFrame", background="#ffffff")
        st.configure("TLabel", background="#ffffff")
        st.configure("TCheckbutton", background="#ffffff")
        st.configure("Buyuk.TButton", font=YAZI_K, padding=(16, 8))
        st.configure("TLabelframe", background="#ffffff")
        st.configure("TLabelframe.Label", background="#ffffff", font=YAZI_K)

        self.adimlar = [
            ("Ay seçimi", self.adim_baslangic, self.topla_baslangic),
            ("Kadro", self.adim_kadro, self.topla_kadro),
            ("Vardiyalar", self.adim_vardiya, self.topla_vardiya),
            ("Performans verileri", self.adim_performans, self.topla_performans),
            ("Makine parkuru", self.adim_makine, self.topla_makine),
            ("Yapılan işler", lambda: self.adim_isler("isler"), self.topla_isler),
            ("Eğitim ve etkinlikler", lambda: self.adim_isler("etkinlikler"), self.topla_isler),
            ("Kontrol ve oluştur", self.adim_son, None),
        ]

        ust = tk.Frame(self.root, bg=MAVI_HEX)
        ust.pack(fill="x")
        tk.Label(ust, text="VEGA TAURUS AVM  •  Temizlik Birimi Aylık Sunum", bg=MAVI_HEX, fg="white",
                 font=("Segoe UI", 14, "bold"), pady=10, padx=16).pack(side="left")
        self.ilerleme = tk.Label(ust, text="", bg=MAVI_HEX, fg="white", font=YAZI, padx=16)
        self.ilerleme.pack(side="right")

        # alt çubuk gövdeden ÖNCE yerleşmeli, yoksa uzun içerikte ekrandan taşar
        alt = ttk.Frame(self.root, style="Beyaz.TFrame", padding=(24, 10))
        alt.pack(fill="x", side="bottom")
        ttk.Separator(self.root).pack(fill="x", side="bottom")
        self.govde = ttk.Frame(self.root, style="Beyaz.TFrame", padding=(24, 12))
        self.govde.pack(fill="both", expand=True)

        self.geri_btn = ttk.Button(alt, text="◀  Geri", style="Buyuk.TButton", command=self.geri)
        self.geri_btn.pack(side="left")
        self.ileri_btn = ttk.Button(alt, text="İleri  ▶", style="Buyuk.TButton", command=self.ileri)
        self.ileri_btn.pack(side="right")
        self.durum = ttk.Label(alt, text="", foreground=GRI)
        self.durum.pack(side="left", padx=20)

        self.root.protocol("WM_DELETE_WINDOW", self.kapat)
        self.goster()
        if guncelleme_notu:
            self.root.after(400, lambda: messagebox.showinfo(
                "Program güncellendi", f"Program yeni sürüme güncellendi (sürüm {surum}).\n\n{guncelleme_notu}"))

    # ---- genel akış -------------------------------------------------------
    def temizle(self):
        for w in self.govde.winfo_children():
            w.destroy()
        self.thumbs = []

    def baslik(self, metin, aciklama=None):
        ttk.Label(self.govde, text=metin, font=BASLIK).pack(anchor="w")
        if aciklama:
            ttk.Label(self.govde, text=aciklama, font=ACIKLAMA, foreground=GRI, wraplength=1080,
                      justify="left").pack(anchor="w", pady=(2, 12))

    def goster(self):
        self.temizle()
        ad, kur, _ = self.adimlar[self.adim]
        self.ilerleme.config(text=f"Adım {self.adim + 1} / {len(self.adimlar)}  —  {ad}")
        self.geri_btn.state(["disabled"] if self.adim == 0 else ["!disabled"])
        self.ileri_btn.config(text="İleri  ▶")
        self.ileri_btn.state(["disabled"] if self.adim == len(self.adimlar) - 1 else ["!disabled"])
        kur()

    def kaydet(self):
        if self.v:
            self.kayit.yaz(self.v)
            self.durum.config(text="Taslak kaydedildi " + datetime.datetime.now().strftime("%H:%M"))

    def topla(self):
        topla = self.adimlar[self.adim][2]
        if topla is None:
            return True
        try:
            sonuc = topla()
        except ValueError as e:
            messagebox.showwarning("Eksik / hatalı bilgi", str(e))
            return False
        if sonuc is False:
            return False
        self.kaydet()
        return True

    def ileri(self):
        if self.topla():
            self.adim += 1
            self.goster()

    def geri(self):
        if self.adimlar[self.adim][2] in (self.topla_isler,):
            if not self.is_editorunu_birak():
                return
        else:
            try:
                if self.adimlar[self.adim][2]:
                    self.adimlar[self.adim][2]()
                    self.kaydet()
            except ValueError:
                pass
        self.adim -= 1
        self.goster()

    def kapat(self):
        if self.v and self.adim > 0:
            try:
                fn = self.adimlar[self.adim][2]
                if fn == self.topla_isler:
                    self.is_editorunu_birak(sessiz=True)
                elif fn:
                    fn()
                self.kaydet()
            except Exception:
                pass
        self.root.destroy()

    # ---- 1. başlangıç -----------------------------------------------------
    def adim_baslangic(self):
        self.baslik("Hoş geldiniz",
                    "Bu program aylık faaliyet sunumunun TEMİZLİK bölümünü sizin için hazırlar. "
                    "Sorulara sırayla cevap verin; her adımda bilgileriniz otomatik kaydedilir, programı kapatıp "
                    "sonra kaldığınız yerden devam edebilirsiniz. Sonunda masaüstüne hazır PowerPoint dosyası oluşur.")
        bugun = datetime.date.today()
        varsayilan = (bugun.replace(day=1) - datetime.timedelta(days=1))
        if self.v:
            yil, ay = self.v["yil"], self.v["ay"]
        else:
            yil, ay = varsayilan.year, varsayilan.month
        f = ttk.Frame(self.govde, style="Beyaz.TFrame")
        f.pack(anchor="w", pady=10)
        ttk.Label(f, text="Hangi ayın sunumunu hazırlıyorsunuz?", font=YAZI_K).grid(row=0, column=0, columnspan=4,
                                                                                   sticky="w", pady=(0, 8))
        self.ay_var = tk.StringVar(value=AYLAR[ay - 1])
        self.yil_var = tk.StringVar(value=str(yil))
        ttk.Combobox(f, textvariable=self.ay_var, values=AYLAR, state="readonly", width=12,
                     font=YAZI).grid(row=1, column=0)
        ttk.Combobox(f, textvariable=self.yil_var, values=[str(y) for y in range(bugun.year - 1, bugun.year + 2)],
                     state="readonly", width=8, font=YAZI).grid(row=1, column=1, padx=8)
        ttk.Label(f, text="Sunumu hazırlayan (ad soyad):", font=YAZI_K).grid(row=2, column=0, columnspan=4,
                                                                           sticky="w", pady=(18, 6))
        self.hazirlayan_var = tk.StringVar(value=(self.v or {}).get("hazirlayan", ""))
        ttk.Entry(f, textvariable=self.hazirlayan_var, width=40, font=YAZI).grid(row=3, column=0, columnspan=4,
                                                                                 sticky="w")

        rehber = ttk.LabelFrame(self.govde, text="  Sunumda dikkat edilecek kurallar  ", padding=14)
        rehber.pack(fill="x", pady=24)
        for satir in [
            "✔  Normal yazım kullanın — TAMAMI BÜYÜK HARF yazmayın (program otomatik düzeltir).",
            "✔  Her işin bir açıklama cümlesi olsun: NEREDE + NE YAPILDI + KAPSAMI (mümkünse sayı ile).",
            "✔  Her iş için en fazla 3 fotoğraf seçin; en iyi, en net olanları seçin (öncesi-sonrası varsa onları).",
            "✔  Benzer işleri tek slaytta birleştirin. Toplam hedef: en fazla 20 slayt.",
            "✔  'Tamamlanmıştır', 'sağlanmıştır' gibi kalıpları her cümlede tekrarlamayın.",
        ]:
            ttk.Label(rehber, text=satir, wraplength=1050, justify="left").pack(anchor="w", pady=2)

    def topla_baslangic(self):
        ay = AYLAR.index(self.ay_var.get()) + 1
        yil = int(self.yil_var.get())
        if not self.hazirlayan_var.get().strip():
            raise ValueError("Lütfen sunumu hazırlayan kişinin adını yazın.")
        if not self.v or (self.v["yil"], self.v["ay"]) != (yil, ay):
            v = self.kayit.oku(yil, ay)
            if v:
                messagebox.showinfo("Kaldığınız yerden devam",
                                    f"{AYLAR[ay - 1]} {yil} için daha önce kaydedilmiş bilgiler bulundu. "
                                    f"Kaldığınız yerden devam ediyorsunuz.")
            else:
                v = bos_veri(self.ayar, yil, ay)
                onceki = self.kayit.onceki_ay(yil, ay)
                if onceki:
                    v["kadro"] = copy.deepcopy(onceki["kadro"])
                    v["vardiya"] = copy.deepcopy(onceki["vardiya"])
                    v["makineler"] = copy.deepcopy(onceki["makineler"])
                    v["performans"]["onceki_sikayet"] = onceki["performans"]["sikayet"]
            self.v = v
        self.v["hazirlayan"] = self.hazirlayan_var.get().strip()

    # ---- 2. kadro ---------------------------------------------------------
    def adim_kadro(self):
        self.baslik("Kadro bilgileri",
                    "Fiili kadro: sözleşmede tanımlı olması gereken personel sayısı.  Gerçekleşen kadro: bu ay "
                    "fiilen çalışan personel sayısı. (Geçen ay girdiyseniz rakamlar dolu gelir, sadece değişeni "
                    "düzeltin.)")
        roller = self.ayar["kadro_rolleri"]
        f = ttk.Frame(self.govde, style="Beyaz.TFrame")
        f.pack(anchor="w", pady=10)
        for c, r in enumerate(roller + ["TOPLAM"], start=1):
            ttk.Label(f, text=r, font=YAZI_K).grid(row=0, column=c, padx=10, pady=6)
        self.kadro_vars = {}
        self.kadro_toplam = {}
        for satir, (anahtar, etiket) in enumerate((("fiili", "Fiili kadro"), ("gerceklesen", "Gerçekleşen kadro")),
                                                  start=1):
            ttk.Label(f, text=etiket, font=YAZI_K).grid(row=satir, column=0, sticky="w", padx=(0, 12), pady=6)
            vars_ = []
            for c, d in enumerate(self.v["kadro"][anahtar], start=1):
                e, var = sayi_kutusu(f, d, degisince=self.kadro_toplamlari)
                e.grid(row=satir, column=c, padx=10, pady=6)
                vars_.append(var)
            self.kadro_vars[anahtar] = vars_
            lbl = ttk.Label(f, text="", font=YAZI_K)
            lbl.grid(row=satir, column=len(roller) + 1)
            self.kadro_toplam[anahtar] = lbl
        self.kadro_uyari = ttk.Label(self.govde, text="", foreground=TURUNCU, font=YAZI_K)
        self.kadro_uyari.pack(anchor="w", pady=10)
        self.kadro_toplamlari()

    def kadro_toplamlari(self):
        toplam = {}
        for k, vars_ in self.kadro_vars.items():
            t = sum(int(v.get()) for v in vars_ if v.get().strip().isdigit())
            toplam[k] = t
            self.kadro_toplam[k].config(text=str(t))
        eksik = toplam.get("fiili", 0) - toplam.get("gerceklesen", 0)
        self.kadro_uyari.config(text=f"Eksik personel: {eksik}" if eksik > 0 else "")

    def topla_kadro(self):
        roller = self.ayar["kadro_rolleri"]
        for k, etiket in (("fiili", "Fiili kadro"), ("gerceklesen", "Gerçekleşen kadro")):
            self.v["kadro"][k] = [sayi_oku(var, f"{etiket} — {r}") for var, r in zip(self.kadro_vars[k], roller)]
        if sum(self.v["kadro"]["fiili"]) == 0:
            raise ValueError("Fiili kadro toplamı 0 olamaz.")

    # ---- 3. vardiya -------------------------------------------------------
    def adim_vardiya(self):
        self.baslik("Vardiya dağılımı",
                    "Her vardiya için planlanan (olması gereken) ve gerçekleşen (fiilen çalışan) personel sayısını "
                    "yazın.")
        roller = self.ayar["vardiya_rolleri"]
        f = ttk.Frame(self.govde, style="Beyaz.TFrame")
        f.pack(anchor="w", pady=10)
        for c, r in enumerate(roller + ["TOPLAM"], start=1):
            ttk.Label(f, text=r, font=YAZI_K).grid(row=0, column=c, padx=12, pady=6)
        self.vardiya_vars = []
        self.vardiya_toplam = []
        satir = 1
        for i, vd in enumerate(self.v["vardiya"], start=1):
            kayit = {}
            for anahtar, etiket in (("plan", f"{i}. Vardiya (planlanan)"), ("gercek", f"{i}. Vardiya (gerçekleşen)")):
                ttk.Label(f, text=etiket, font=YAZI_K if anahtar == "plan" else YAZI).grid(
                    row=satir, column=0, sticky="w", padx=(0, 12), pady=5)
                vars_ = []
                for c, d in enumerate(vd[anahtar], start=1):
                    e, var = sayi_kutusu(f, d, degisince=self.vardiya_toplamlari)
                    e.grid(row=satir, column=c, padx=12, pady=5)
                    vars_.append(var)
                lbl = ttk.Label(f, text="", font=YAZI_K)
                lbl.grid(row=satir, column=len(roller) + 1)
                kayit[anahtar] = (vars_, lbl)
                satir += 1
            self.vardiya_vars.append(kayit)
            ttk.Separator(f).grid(row=satir, column=0, columnspan=len(roller) + 2, sticky="ew", pady=4)
            satir += 1
        self.vardiya_bilgi = ttk.Label(self.govde, text="", foreground=GRI)
        self.vardiya_bilgi.pack(anchor="w", pady=8)
        self.vardiya_toplamlari()

    def vardiya_toplamlari(self):
        gercek = 0
        for kayit in self.vardiya_vars:
            for anahtar, (vars_, lbl) in kayit.items():
                t = sum(int(v.get()) for v in vars_ if v.get().strip().isdigit())
                lbl.config(text=str(t))
                if anahtar == "gercek":
                    gercek += t
        kadro = sum(self.v["kadro"]["gerceklesen"])
        self.vardiya_bilgi.config(
            text=f"Vardiyalarda gerçekleşen toplam: {gercek}   •   Önceki adımda girilen gerçekleşen kadro: {kadro}",
            foreground=TURUNCU if gercek != kadro else GRI)

    def topla_vardiya(self):
        roller = self.ayar["vardiya_rolleri"]
        for i, kayit in enumerate(self.vardiya_vars):
            for anahtar, ad in (("plan", "planlanan"), ("gercek", "gerçekleşen")):
                self.v["vardiya"][i][anahtar] = [sayi_oku(var, f"{i + 1}. vardiya {ad} — {r}")
                                                 for var, r in zip(kayit[anahtar][0], roller)]
        gercek = sum(sum(vd["gercek"]) for vd in self.v["vardiya"])
        kadro = sum(self.v["kadro"]["gerceklesen"])
        if gercek != kadro:
            return messagebox.askyesno(
                "Rakamlar tutmuyor",
                f"Vardiyalarda gerçekleşen personel toplamı {gercek}, ama gerçekleşen kadro {kadro} girildi.\n\n"
                "Rakamlar doğru mu? (Evet = yine de devam et, Hayır = düzelteyim)")
        return True

    # ---- 4. performans ----------------------------------------------------
    def adim_performans(self):
        self.baslik("Performans verileri",
                    "Üst yönetim 'yapıldı' değil 'ne kadar, ne değişti' görmek ister. Bu rakamlar sunumdaki Genel "
                    "Durum slaytına otomatik yazılır ve geçen ayla karşılaştırılır.")
        p = self.v["performans"]
        f = ttk.Frame(self.govde, style="Beyaz.TFrame")
        f.pack(anchor="w", pady=6)
        self.perf_vars = {}
        alanlar = [
            ("sikayet", "Bu ay gelen şikayet sayısı (ziyaretçi + mağaza + yönetim)"),
            ("onceki_sikayet", "Geçen ay gelen şikayet sayısı"),
            ("denetim", "Bu ay yapılan hijyen denetimi sayısı"),
            ("uygunsuzluk", "Denetimlerde tespit edilen uygunsuzluk sayısı"),
            ("egitim", "Bu ay eğitim verilen personel sayısı (yoksa 0)"),
        ]
        for satir, (k, etiket) in enumerate(alanlar):
            ttk.Label(f, text=etiket).grid(row=satir, column=0, sticky="w", pady=7, padx=(0, 16))
            e, var = sayi_kutusu(f, p.get(k), genislik=8)
            e.grid(row=satir, column=1, sticky="w")
            self.perf_vars[k] = var
        if p.get("onceki_sikayet") is not None:
            ttk.Label(f, text="(geçen ayki kayıttan otomatik geldi)", foreground=GRI, font=ACIKLAMA).grid(
                row=1, column=2, sticky="w", padx=8)
        ttk.Label(self.govde, text="Eklemek istediğiniz kısa değerlendirme (isteğe bağlı, 1-2 cümle):",
                  font=YAZI_K).pack(anchor="w", pady=(16, 4))
        self.ek_not = tk.Text(self.govde, height=3, width=110, font=YAZI, wrap="word", relief="solid", bd=1)
        self.ek_not.insert("1.0", p.get("ek_not", ""))
        self.ek_not.pack(anchor="w")

    def topla_performans(self):
        p = self.v["performans"]
        for k, ad in (("sikayet", "Bu ay şikayet"), ("onceki_sikayet", "Geçen ay şikayet"),
                      ("denetim", "Hijyen denetimi"), ("uygunsuzluk", "Uygunsuzluk"), ("egitim", "Eğitim")):
            p[k] = sayi_oku(self.perf_vars[k], ad)
        if p["uygunsuzluk"] and not p["denetim"]:
            raise ValueError("Uygunsuzluk sayısı girilmiş ama denetim sayısı 0. Lütfen kontrol edin.")
        ek = self.ek_not.get("1.0", "end").strip()
        p["ek_not"] = metni_duzelt(ek, self.ayar)[0] if ek else ""

    # ---- 5. makine --------------------------------------------------------
    def adim_makine(self):
        self.baslik("Makine dairesi ve makine parkuru",
                    "Makine dairesinin bu ayki durumunu gösteren en fazla 2 fotoğraf seçin. Altta makine listesini "
                    "güncelleyin: durumu değişen makinenin 'Durum' kutusunu değiştirin, yeni makine geldiyse ekleyin.")
        ust = ttk.LabelFrame(self.govde, text="  Makine dairesi fotoğrafları (en fazla 2)  ", padding=10)
        ust.pack(fill="x")
        self.md_foto_alani = ttk.Frame(ust, style="Beyaz.TFrame")
        self.md_foto_alani.pack(side="left", fill="x", expand=True)
        sag = ttk.Frame(ust, style="Beyaz.TFrame")
        sag.pack(side="right")
        ttk.Button(sag, text="📷  Fotoğraf seç", command=self.md_foto_ekle).pack(anchor="e")
        ttk.Label(sag, text="Kısa açıklama (isteğe bağlı):", font=ACIKLAMA).pack(anchor="e", pady=(8, 0))
        self.md_not = tk.StringVar(value=self.v["makine_dairesi"].get("not", ""))
        ttk.Entry(sag, textvariable=self.md_not, width=45).pack(anchor="e")
        self.md_fotolari_ciz()

        alt = ttk.LabelFrame(self.govde, text="  Makine parkuru  ", padding=8)
        alt.pack(fill="both", expand=True, pady=(12, 0))
        bas = ttk.Frame(alt, style="Beyaz.TFrame")
        bas.pack(fill="x")
        for c, (t, w) in enumerate((("No", 4), ("Marka", 20), ("Model", 12), ("Makine özelliği", 34),
                                    ("Durum", 12), ("", 6))):
            ttk.Label(bas, text=t, font=YAZI_K, width=w).grid(row=0, column=c, padx=3, sticky="w")
        self.makine_kaydir = Kaydirmali(alt)
        self.makine_kaydir.pack(fill="both", expand=True)
        ttk.Button(alt, text="＋  Makine ekle", command=self.makine_ekle).pack(anchor="w", pady=(6, 0))
        self.makine_satirlari = []
        for m in self.v["makineler"]:
            self.makine_satiri(m)

    def md_fotolari_ciz(self):
        for w in self.md_foto_alani.winfo_children():
            w.destroy()
        fotolar = self.v["makine_dairesi"]["fotolar"]
        if not fotolar:
            ttk.Label(self.md_foto_alani, text="Henüz fotoğraf seçilmedi.", foreground=GRI).pack(anchor="w")
        for i, f in enumerate(fotolar):
            self.kucuk_resim(self.md_foto_alani, f, lambda i=i: self.md_foto_sil(i)).pack(side="left", padx=6)

    def md_foto_ekle(self):
        fotolar = self.v["makine_dairesi"]["fotolar"]
        kalan = self.ayar["makine_dairesi_max_foto"] - len(fotolar)
        if kalan <= 0:
            messagebox.showinfo("Sınır", "En fazla 2 fotoğraf seçebilirsiniz. Önce birini kaldırın.")
            return
        fotolar.extend(self.foto_sec(kalan))
        self.md_fotolari_ciz()
        self.kaydet()

    def md_foto_sil(self, i):
        del self.v["makine_dairesi"]["fotolar"][i]
        self.md_fotolari_ciz()
        self.kaydet()

    def makine_satiri(self, m):
        f = ttk.Frame(self.makine_kaydir.ic, style="Beyaz.TFrame")
        f.pack(fill="x", pady=2)
        no = ttk.Label(f, text="", width=4)
        no.grid(row=0, column=0, padx=3)
        vars_ = [tk.StringVar(value=x) for x in m]
        ttk.Entry(f, textvariable=vars_[0], width=20).grid(row=0, column=1, padx=3)
        ttk.Entry(f, textvariable=vars_[1], width=12).grid(row=0, column=2, padx=3)
        ttk.Entry(f, textvariable=vars_[2], width=34).grid(row=0, column=3, padx=3)
        ttk.Combobox(f, textvariable=vars_[3], values=self.ayar["makine_durumlari"], width=11,
                     state="readonly").grid(row=0, column=4, padx=3)
        satir = {"frame": f, "vars": vars_, "no": no}
        ttk.Button(f, text="Sil", width=5, command=lambda: self.makine_sil(satir)).grid(row=0, column=5, padx=3)
        self.makine_satirlari.append(satir)
        self.makine_numarala()

    def makine_ekle(self):
        self.makine_satiri(["", "", "", "AKTİF"])
        self.root.after(50, lambda: self.makine_kaydir.canvas.yview_moveto(1.0))

    def makine_sil(self, satir):
        satir["frame"].destroy()
        self.makine_satirlari.remove(satir)
        self.makine_numarala()

    def makine_numarala(self):
        for i, s in enumerate(self.makine_satirlari, start=1):
            s["no"].config(text=str(i))

    def topla_makine(self):
        liste = []
        for i, s in enumerate(self.makine_satirlari, start=1):
            m = [v.get().strip() for v in s["vars"]]
            if not any(m[:3]):
                continue
            if not m[0] or not m[2]:
                raise ValueError(f"{i}. makinenin marka veya özellik bilgisi boş.")
            m[0], m[2] = tr_upper(m[0]), tr_upper(m[2])
            liste.append(m)
        if not liste:
            raise ValueError("Makine listesi boş olamaz.")
        self.v["makineler"] = liste
        self.v["makine_dairesi"]["not"] = metni_duzelt(self.md_not.get(), self.ayar)[0] if self.md_not.get().strip() else ""
        if not self.v["makine_dairesi"]["fotolar"]:
            return messagebox.askyesno("Fotoğraf yok",
                                       "Makine dairesi fotoğrafı seçmediniz. Bu slayt sunuma eklenmeyecek.\n\n"
                                       "Fotoğrafsız devam edilsin mi?")
        return True

    # ---- fotoğraf yardımcıları --------------------------------------------
    def foto_sec(self, en_fazla):
        yollar = filedialog.askopenfilenames(
            title=f"Fotoğraf seçin (en fazla {en_fazla} adet)",
            filetypes=[("Fotoğraflar", "*.jpg *.jpeg *.png *.bmp *.webp *.JPG *.JPEG *.PNG"), ("Tüm dosyalar", "*.*")])
        if not yollar:
            return []
        if len(yollar) > en_fazla:
            messagebox.showinfo("Fazla fotoğraf",
                                f"{len(yollar)} fotoğraf seçtiniz, en fazla {en_fazla} tane eklenebilir. "
                                f"İlk {en_fazla} tanesi alındı. Sunumda en iyi fotoğrafların olması önemli — "
                                "isterseniz kaldırıp başkasını seçebilirsiniz.")
            yollar = yollar[:en_fazla]
        sonuc = []
        hedef = os.path.join(self.kayit.klasor(self.v["yil"], self.v["ay"]), "foto")
        self.root.config(cursor="watch")
        self.root.update()
        try:
            for y in yollar:
                try:
                    sonuc.append(fotografi_hazirla(y, hedef))
                except Exception:
                    messagebox.showerror("Fotoğraf açılamadı",
                                         f"Bu dosya açılamadı:\n{os.path.basename(y)}\n\n"
                                         "iPhone (HEIC) fotoğrafıysa WhatsApp'tan bilgisayara gönderip öyle seçin.")
        finally:
            self.root.config(cursor="")
        return sonuc

    def kucuk_resim(self, parent, f, sil_fonk, boyut=130):
        kutu = ttk.Frame(parent, style="Beyaz.TFrame")
        try:
            im = Image.open(os.path.join(self.kayit.klasor(self.v["yil"], self.v["ay"]), "foto", f["dosya"]))
            im.thumbnail((boyut, boyut))
            ph = ImageTk.PhotoImage(im)
            self.thumbs.append(ph)
            tk.Label(kutu, image=ph, bg="#ffffff", relief="solid", bd=1).pack()
        except Exception:
            ttk.Label(kutu, text="(fotoğraf bulunamadı)", foreground=KIRMIZI_HEX).pack()
        ttk.Button(kutu, text="✕ Kaldır", command=sil_fonk).pack(pady=2)
        return kutu

    # ---- 6-7. işler / etkinlikler -----------------------------------------
    def adim_isler(self, kategori):
        self.kategori = kategori
        if kategori == "isler":
            self.baslik("Bu ay yapılan işler",
                        "Her iş bir slayt olur. Sağda cümleyi yazın, 1-3 fotoğraf seçin, 'Listeye kaydet'e basın. "
                        "Benzer işleri (ör. tüm otopark yıkamaları) tek işte birleştirin. Sırayı ▲▼ ile "
                        "değiştirebilirsiniz — sunumdaki sıra listedeki sıradır.")
            ornekler = self.ayar["ornek_cumleler"]
        else:
            self.baslik("Eğitim ve etkinlikler (isteğe bağlı)",
                        "Personel eğitimleri, ayın elemanı, doğum günü vb. Her biri bir slayt olur. Yoksa bu adımı "
                        "boş geçip 'İleri'ye basın. Kural aynı: kişi sayısı / konu / tarih gibi somut bilgi yazın.")
            ornekler = self.ayar["etkinlik_ornekleri"]
        self.liste = self.v[kategori]
        self.secili = None
        self.duzenleme_fotolari = []

        ana = ttk.Frame(self.govde, style="Beyaz.TFrame")
        ana.pack(fill="both", expand=True)
        sol = ttk.Frame(ana, style="Beyaz.TFrame")
        sol.pack(side="left", fill="y")
        ttk.Label(sol, text="Liste", font=YAZI_K).pack(anchor="w")
        lf = ttk.Frame(sol, style="Beyaz.TFrame")
        lf.pack(fill="y", expand=True)
        self.lb = tk.Listbox(lf, width=46, font=("Segoe UI", 10), activestyle="none", selectbackground=MAVI_HEX,
                             exportselection=False)
        sb = ttk.Scrollbar(lf, command=self.lb.yview)
        self.lb.configure(yscrollcommand=sb.set)
        self.lb.pack(side="left", fill="y", expand=True)
        sb.pack(side="left", fill="y")
        self.lb.bind("<<ListboxSelect>>", self.listeden_sec)
        bf = ttk.Frame(sol, style="Beyaz.TFrame")
        bf.pack(fill="x", pady=6)
        ttk.Button(bf, text="＋ Yeni", command=self.yeni_is).pack(side="left")
        ttk.Button(bf, text="▲", width=3, command=lambda: self.tasi(-1)).pack(side="left", padx=4)
        ttk.Button(bf, text="▼", width=3, command=lambda: self.tasi(1)).pack(side="left")
        ttk.Button(bf, text="Sil", command=self.is_sil).pack(side="right")
        self.sayac = ttk.Label(sol, text="", foreground=GRI)
        self.sayac.pack(anchor="w")

        sag = ttk.Frame(ana, style="Beyaz.TFrame", padding=(20, 0, 0, 0))
        sag.pack(side="left", fill="both", expand=True)
        self.editor_baslik = ttk.Label(sag, text="Yeni iş", font=YAZI_K)
        self.editor_baslik.pack(anchor="w")
        ttk.Label(sag, text="Açıklama cümlesi:", font=ACIKLAMA, foreground=GRI).pack(anchor="w", pady=(6, 2))
        self.metin = tk.Text(sag, height=3, font=("Segoe UI", 12), wrap="word", relief="solid", bd=1)
        self.metin.pack(fill="x")
        self.metin.bind("<KeyRelease>", lambda e: self.canli_kontrol())
        ornek = ttk.Label(sag, text="Örnek: " + "\nÖrnek: ".join(ornekler[:2]), font=ACIKLAMA, foreground=GRI,
                          wraplength=650, justify="left")
        ornek.pack(anchor="w", pady=(4, 0))
        self.uyari_lbl = ttk.Label(sag, text="", foreground=TURUNCU, wraplength=650, justify="left")
        self.uyari_lbl.pack(anchor="w", pady=(6, 0))
        ff = ttk.Frame(sag, style="Beyaz.TFrame")
        ff.pack(fill="x", pady=(10, 0))
        ttk.Label(ff, text=f"Fotoğraflar (en az 1, en fazla {self.ayar['max_foto']}):", font=YAZI_K).pack(side="left")
        ttk.Button(ff, text="📷  Fotoğraf seç", command=self.is_foto_ekle).pack(side="left", padx=10)
        ttk.Button(ff, text="✔  Listeye kaydet", style="Buyuk.TButton", command=self.is_kaydet).pack(side="right")
        self.foto_alani = ttk.Frame(sag, style="Beyaz.TFrame")
        self.foto_alani.pack(fill="x", pady=6)
        self.listeyi_ciz()
        self.editoru_ac(None)

    def listeyi_ciz(self):
        self.lb.delete(0, "end")
        for i, is_ in enumerate(self.liste, start=1):
            m = is_["metin"]
            self.lb.insert("end", f"{i}. {m[:52]}{'…' if len(m) > 52 else ''}  [{len(is_['fotolar'])} foto]")
        if self.secili is not None and self.secili < len(self.liste):
            self.lb.selection_set(self.secili)
        tum = len(self.v["isler"]) + len(self.v["etkinlikler"])
        self.sayac.config(text=f"Bu listede {len(self.liste)} slayt  •  Toplam iş/etkinlik slaytı: {tum}")

    def editoru_ac(self, i):
        self.secili = i
        self.metin.delete("1.0", "end")
        if i is None:
            self.editor_baslik.config(text="Yeni iş" if self.kategori == "isler" else "Yeni etkinlik / eğitim")
            self.duzenleme_fotolari = []
            self.lb.selection_clear(0, "end")
        else:
            self.editor_baslik.config(text=f"{i + 1}. kaydı düzenliyorsunuz")
            self.metin.insert("1.0", self.liste[i]["metin"])
            self.duzenleme_fotolari = list(self.liste[i]["fotolar"])
        self.editor_ilk = (self.metin.get("1.0", "end").strip(), [f["dosya"] for f in self.duzenleme_fotolari])
        self.fotolari_ciz()
        self.canli_kontrol()

    def editor_degisti(self):
        su_an = (self.metin.get("1.0", "end").strip(), [f["dosya"] for f in self.duzenleme_fotolari])
        return su_an != self.editor_ilk

    def is_editorunu_birak(self, sessiz=False):
        """Editörde kaydedilmemiş değişiklik varsa sorar. Devam edilebilirse True."""
        if not self.editor_degisti():
            return True
        if sessiz:
            return self.is_kaydet(sessiz=True) or True
        cevap = messagebox.askyesnocancel("Kaydedilmemiş değişiklik",
                                          "Sağdaki kutuda listeye kaydedilmemiş bir iş var.\n\n"
                                          "Evet = listeye kaydet\nHayır = vazgeç (sil)\nİptal = bu ekranda kal")
        if cevap is None:
            return False
        if cevap:
            return self.is_kaydet()
        return True

    def listeden_sec(self, _=None):
        sec = self.lb.curselection()
        if not sec or sec[0] == self.secili:
            return
        if not self.is_editorunu_birak():
            self.listeyi_ciz()
            return
        self.editoru_ac(sec[0])

    def yeni_is(self):
        if self.is_editorunu_birak():
            self.editoru_ac(None)
            self.metin.focus_set()

    def tasi(self, yon):
        i = self.secili
        if i is None:
            return
        j = i + yon
        if 0 <= j < len(self.liste):
            self.liste[i], self.liste[j] = self.liste[j], self.liste[i]
            self.secili = j
            self.listeyi_ciz()
            self.editor_baslik.config(text=f"{j + 1}. kaydı düzenliyorsunuz")
            self.kaydet()

    def is_sil(self):
        if self.secili is None:
            return
        if messagebox.askyesno("Sil", f"{self.secili + 1}. kayıt silinsin mi?"):
            del self.liste[self.secili]
            self.kaydet()
            self.editoru_ac(None)
            self.listeyi_ciz()

    def fotolari_ciz(self):
        for w in self.foto_alani.winfo_children():
            w.destroy()
        self.thumbs = []
        if not self.duzenleme_fotolari:
            ttk.Label(self.foto_alani, text="Henüz fotoğraf seçilmedi.", foreground=GRI).pack(anchor="w")
        for i, f in enumerate(self.duzenleme_fotolari):
            self.kucuk_resim(self.foto_alani, f, lambda i=i: self.is_foto_sil(i), boyut=120).pack(side="left", padx=6)

    def is_foto_ekle(self):
        kalan = self.ayar["max_foto"] - len(self.duzenleme_fotolari)
        if kalan <= 0:
            messagebox.showinfo("Sınır", f"Bir slaytta en fazla {self.ayar['max_foto']} fotoğraf olabilir. "
                                         "Daha fazla fotoğrafınız varsa en iyilerini seçin.")
            return
        self.duzenleme_fotolari.extend(self.foto_sec(kalan))
        self.fotolari_ciz()
        self.canli_kontrol()

    def is_foto_sil(self, i):
        del self.duzenleme_fotolari[i]
        self.fotolari_ciz()

    def canli_kontrol(self):
        t = self.metin.get("1.0", "end").strip()
        if not t:
            self.uyari_lbl.config(text="")
            return
        duz, _ = metni_duzelt(t, self.ayar)
        uyarilar = metin_uyarilari(duz, self.ayar)
        self.uyari_lbl.config(text="\n".join("⚠ " + u for u in uyarilar) if uyarilar else "✔ Cümle uygun görünüyor.",
                              foreground=TURUNCU if uyarilar else YESIL)

    def is_kaydet(self, sessiz=False):
        t = self.metin.get("1.0", "end").strip()
        if not t:
            if not sessiz:
                messagebox.showwarning("Eksik", "Açıklama cümlesi boş.")
            return False
        if not self.duzenleme_fotolari:
            if not sessiz:
                messagebox.showwarning("Eksik", "En az 1 fotoğraf seçmelisiniz. Açıklamasız fotoğraf olmadığı gibi "
                                                "fotoğrafsız açıklama da sunuma girmez.")
            return False
        duz, degisiklik = metni_duzelt(t, self.ayar)
        uyarilar = metin_uyarilari(duz, self.ayar)
        ciddi = [u for u in uyarilar if not u.startswith("İpucu")]
        if ciddi and not sessiz:
            if not messagebox.askyesno("Cümleyi kontrol edin",
                                       "\n\n".join("⚠ " + u for u in ciddi) +
                                       "\n\nYine de bu haliyle kaydedilsin mi?\n(Hayır = geri dönüp düzelteyim)"):
                return False
        kayit = {"metin": duz, "fotolar": list(self.duzenleme_fotolari)}
        if self.secili is None:
            self.liste.append(kayit)
            self.secili = len(self.liste) - 1
        else:
            self.liste[self.secili] = kayit
        self.kaydet()
        if degisiklik and not sessiz:
            messagebox.showinfo("Otomatik düzeltme", "Cümleniz düzeltildi:\n\n" + duz + "\n\n• " + "\n• ".join(
                degisiklik))
        self.listeyi_ciz()
        self.editoru_ac(None)
        return True

    def topla_isler(self):
        if not self.is_editorunu_birak():
            return False
        if self.kategori == "isler" and not self.v["isler"]:
            raise ValueError("En az bir yapılan iş eklemelisiniz.")
        return True

    # ---- 8. kontrol ve oluştur --------------------------------------------
    def kontroller(self):
        v = self.v
        sonuc = []  # (durum, metin)  durum: ok / uyari / hata
        fiili = sum(v["kadro"]["fiili"])
        sonuc.append(("ok" if fiili else "hata", f"Kadro bilgisi girildi (tanımlı {fiili}, "
                                                 f"fiilen çalışan {sum(v['kadro']['gerceklesen'])})."))
        sonuc.append(("ok", f"Performans: {v['performans']['sikayet']} şikayet, {v['performans']['denetim']} "
                            f"hijyen denetimi."))
        tum = v["isler"] + v["etkinlikler"]
        fotosuz = [i + 1 for i, x in enumerate(tum) if not x["fotolar"]]
        sonuc.append(("hata" if fotosuz else "ok", "Her işin fotoğrafı ve açıklaması var." if not fotosuz
                      else f"Fotoğrafsız kayıt var: {fotosuz}"))
        fazla = [i + 1 for i, x in enumerate(tum) if len(x["fotolar"]) > self.ayar["max_foto"]]
        sonuc.append(("hata" if fazla else "ok", f"Slayt başına en fazla {self.ayar['max_foto']} fotoğraf."))
        slayt = 4 + (1 if v["makine_dairesi"]["fotolar"] else 0) + len(tum)
        hedef = self.ayar["hedef_max_slayt"]
        sonuc.append(("uyari" if slayt > hedef else "ok",
                      f"Toplam {slayt} slayt" + (f" — hedef en fazla {hedef}. Benzer işleri birleştirmeyi "
                                                 "düşünün." if slayt > hedef else f" (hedef ≤ {hedef}).")))
        tekrar = {}
        for x in tum:
            dusuk = tr_lower(x["metin"])
            for k in self.ayar["kalip_ifadeler"]:
                if k in dusuk:
                    tekrar[k] = tekrar.get(k, 0) + 1
        cok = {k: n for k, n in tekrar.items() if n > self.ayar["kalip_max_tekrar"]}
        sonuc.append(("uyari" if cok else "ok",
                      "Kalıp ifade tekrarı yok." if not cok else "Kalıp ifade çok tekrarlanıyor: " + ", ".join(
                          f"'{k}' {n} kez" for k, n in cok.items()) + ". Cümleleri çeşitlendirin (ör. 'yapılmıştır', "
                                                                      "'temizlenmiştir', 'yenilenmiştir')."))
        kisa = [i + 1 for i, x in enumerate(tum) if len(x["metin"].split()) < self.ayar["min_kelime"]]
        sonuc.append(("uyari" if kisa else "ok", "Açıklamalar yeterince detaylı." if not kisa
                      else f"Kısa açıklamalar (sıra no): {kisa} — konum ve kapsam ekleyin."))
        buyuk = [i + 1 for i, x in enumerate(tum) if x["metin"] != metni_duzelt(x["metin"], self.ayar)[0]]
        sonuc.append(("uyari" if buyuk else "ok", "Yazım biçimi standarda uygun (BÜYÜK HARF yok)." if not buyuk
                      else f"Biçimi düzeltilecek kayıtlar: {buyuk} (oluştururken otomatik düzeltilir)."))
        return sonuc, slayt

    def adim_son(self):
        self.baslik("Kontrol ve sunumu oluştur",
                    "Aşağıdaki kontrol listesi otomatik hazırlandı. Kırmızı madde varsa geri dönüp düzeltin; turuncu "
                    "maddeler öneridir. Hazırsanız 'Sunumu oluştur'a basın.")
        sonuc, slayt = self.kontroller()
        kf = ttk.LabelFrame(self.govde, text="  Kontrol listesi  ", padding=12)
        kf.pack(fill="x")
        isaret = {"ok": ("✔", YESIL), "uyari": ("⚠", TURUNCU), "hata": ("✖", KIRMIZI_HEX)}
        for durum, metin in sonuc:
            s, renk = isaret[durum]
            ttk.Label(kf, text=f"{s}  {metin}", foreground=renk, wraplength=1050, justify="left").pack(anchor="w",
                                                                                                    pady=2)
        self.hata_var = any(d == "hata" for d, _ in sonuc)

        of = ttk.LabelFrame(self.govde, text="  Genel Durum slaytına yazılacak metin (otomatik)  ", padding=12)
        of.pack(fill="x", pady=12)
        ttk.Label(of, text=genel_durum_cumlesi(self.v), wraplength=1050, justify="left").pack(anchor="w")

        bf = ttk.Frame(self.govde, style="Beyaz.TFrame")
        bf.pack(fill="x", pady=8)
        b = ttk.Button(bf, text="📊  Sunumu oluştur", style="Buyuk.TButton", command=self.olustur)
        b.pack(side="left")
        if self.hata_var:
            b.state(["disabled"])
        self.sonuc_lbl = ttk.Label(bf, text="", foreground=YESIL, font=YAZI_K, wraplength=800, justify="left")
        self.sonuc_lbl.pack(side="left", padx=16)

    def olustur(self):
        v = copy.deepcopy(self.v)
        for x in v["isler"] + v["etkinlikler"]:
            x["metin"] = metni_duzelt(x["metin"], self.ayar)[0]
        klasor = os.path.join(masaustu(), "Temizlik Sunumları")
        os.makedirs(klasor, exist_ok=True)
        ad = f"VEGA TAURUS TEMİZLİK - {AYLAR[v['ay'] - 1]} {v['yil']}.pptx"
        yol = os.path.join(klasor, ad)
        self.root.config(cursor="watch")
        self.sonuc_lbl.config(text="Oluşturuluyor, lütfen bekleyin…", foreground=GRI)
        self.root.update()
        try:
            n = sunumu_olustur(v, self.kayit.klasor(v["yil"], v["ay"]), yol)
        except PermissionError:
            messagebox.showerror("Dosya açık", f"'{ad}' şu an PowerPoint'te açık görünüyor.\n\n"
                                               "PowerPoint'i kapatıp tekrar 'Sunumu oluştur'a basın.")
            self.sonuc_lbl.config(text="")
            return
        except Exception:
            hata = traceback.format_exc()
            with open(os.path.join(self.kok, "hata_kaydi.txt"), "a", encoding="utf-8") as f:
                f.write(f"\n=== {datetime.datetime.now()} ===\n{hata}")
            messagebox.showerror("Hata", "Sunum oluşturulurken bir hata oldu. Program klasöründeki "
                                         "'hata_kaydi.txt' dosyasını Mert Bey'e gönderin.\n\n" + hata[-600:])
            self.sonuc_lbl.config(text="")
            return
        finally:
            self.root.config(cursor="")
        self.v["tamamlandi"] = True
        self.kaydet()
        self.sonuc_lbl.config(text=f"✔ {n} slaytlık sunum hazır:\nMasaüstü › Temizlik Sunumları › {ad}",
                              foreground=YESIL)
        if messagebox.askyesno("Sunum hazır", f"{n} slaytlık sunum oluşturuldu.\n\n{yol}\n\nŞimdi açılsın mı?"):
            os.startfile(yol)
        else:
            subprocess.Popen(["explorer", "/select,", yol])

    def calistir(self):
        self.root.mainloop()


MAVI_HEX = "#4472C4"
KIRMIZI_HEX = "#C00000"


def main(kok, surum="?", guncelleme_notu=None):
    Uygulama(kok, surum, guncelleme_notu).calistir()


if __name__ == "__main__":
    main(os.path.dirname(UYGULAMA_DIR))
